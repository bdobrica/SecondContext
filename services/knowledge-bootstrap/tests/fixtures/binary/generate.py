"""Regenerate small, synthetic fixtures using the locked parser dependencies."""

from pathlib import Path

from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject, NumberObject

ROOT = Path(__file__).parent


def pdf_page(writer, first: str, second: str = ""):
    page = writer.add_blank_page(width=612, height=792)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    page[NameObject("/Resources")] = DictionaryObject(
        {NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})}
    )
    stream = DecodedStreamObject()
    stream.set_data(
        (
            "BT /F1 12 Tf 1 0 0 1 50 730 Tm (" + first + ") Tj "
            "1 0 0 1 50 675 Tm (" + second + ") Tj ET"
        ).encode()
    )
    page[NameObject("/Contents")] = writer._add_object(stream)


def main():
    writer = PdfWriter()
    pdf_page(writer, "Deployment handbook: use a staged rollout.", "Check health before promotion.")
    writer.add_blank_page(width=612, height=792)
    pdf_page(
        writer, "Recovery procedure: roll back the release.", "Keep an audit trail of changes."
    )
    writer.write(ROOT / "handbook.pdf")

    scanned = PdfWriter()
    page = scanned.add_blank_page(width=612, height=792)
    image = DecodedStreamObject()
    image.set_data(b"\x00\xff\xff\x00")
    image.update(
        {
            NameObject("/Type"): NameObject("/XObject"),
            NameObject("/Subtype"): NameObject("/Image"),
            NameObject("/Width"): NumberObject(2),
            NameObject("/Height"): NumberObject(2),
            NameObject("/ColorSpace"): NameObject("/DeviceGray"),
            NameObject("/BitsPerComponent"): NumberObject(8),
        }
    )
    page[NameObject("/Resources")] = DictionaryObject(
        {
            NameObject("/XObject"): DictionaryObject(
                {NameObject("/Scan"): scanned._add_object(image)}
            )
        }
    )
    stream = DecodedStreamObject()
    stream.set_data(b"q 500 0 0 700 50 50 cm /Scan Do Q")
    page[NameObject("/Contents")] = scanned._add_object(stream)
    scanned.write(ROOT / "image-only.pdf")

    document = Document()
    document.add_heading("Deployment handbook", level=1)
    document.add_paragraph("Deploy safely. Preserve café and Romanian diacritics: țară.")
    document.add_heading("Checks", level=2)
    document.add_paragraph("Check health", style="List Bullet")
    document.add_paragraph("Check logs", style="List Bullet")
    document.add_paragraph("Stage release", style="List Number")
    document.add_paragraph("Promote release", style="List Number")
    table = document.add_table(rows=3, cols=2)
    for row, values in zip(
        table.rows,
        [("Environment", "Policy"), ("staging", "automatic"), ("production", "manual")],
        strict=True,
    ):
        for cell, value in zip(row.cells, values, strict=True):
            cell.text = value
    document.add_heading("Recovery", level=1)
    document.add_paragraph("Roll back; keep an audit trail.")
    document.save(ROOT / "handbook.docx")


if __name__ == "__main__":
    main()
