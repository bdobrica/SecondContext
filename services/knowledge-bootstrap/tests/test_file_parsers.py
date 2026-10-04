import json
import subprocess
import sys
from io import BytesIO
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from pypdf import PdfReader, PdfWriter
from pypdf.generic import DecodedStreamObject, NameObject

from knowledge_bootstrap.binary import parse_binary, select_upload_format
from knowledge_bootstrap.parsers import ParseError

FIXTURES = Path(__file__).parent / "fixtures" / "binary"


def binary_fixture(fmt):
    return (FIXTURES / f"handbook.{fmt}").read_bytes()


def repack_docx(changes=None, additions=None):
    output = BytesIO()
    with (
        ZipFile(BytesIO(binary_fixture("docx"))) as original,
        ZipFile(output, "w", ZIP_DEFLATED) as target,
    ):
        for member in original.infolist():
            content = original.read(member)
            if changes and member.filename in changes:
                content = changes[member.filename](content)
            target.writestr(member.filename, content)
        for name, value in (additions or {}).items():
            target.writestr(name, value)
    return output.getvalue()


def assert_error(data, fmt, settings, code):
    with pytest.raises(ParseError) as caught:
        parse_binary(data, fmt, settings)
    assert caught.value.code == code
    assert "SECRET" not in caught.value.detail


def test_pdf_pages_paragraphs_and_determinism(settings):
    first = parse_binary(binary_fixture("pdf"), "pdf", settings)
    second = parse_binary(binary_fixture("pdf"), "pdf", settings)
    assert first == second
    assert first.content_hash == second.content_hash
    assert first.extra_metadata == {"page_count": 3, "pages_with_low_text": [2]}
    assert [block.page_start for block in first.blocks] == [1, 1, 3, 3]
    assert all(block.page_start == block.page_end for block in first.blocks)
    assert first.blocks[0].text == "Deployment handbook: use a staged rollout."
    assert first.blocks[1].text == "Check health before promotion."
    assert "\n\n" in first.text


def test_docx_headings_lists_tables_unicode_and_determinism(settings):
    first = parse_binary(binary_fixture("docx"), "docx", settings)
    assert first == parse_binary(binary_fixture("docx"), "docx", settings)
    assert first.title == "Deployment handbook"
    assert "café" in first.text and "țară" in first.text
    lists = [block for block in first.blocks if block.type == "list"]
    assert [block.text for block in lists] == [
        "- Check health",
        "- Check logs",
        "1. Stage release",
        "2. Promote release",
    ]
    assert all(block.heading_path == ["Deployment handbook", "Checks"] for block in lists)
    table = next(block for block in first.blocks if block.type == "table")
    assert table.text == "Environment | Policy\nstaging | automatic\nproduction | manual"
    assert table.heading_path == ["Deployment handbook", "Checks"]
    assert first.blocks[-1].heading_path == ["Recovery"]
    assert first.blocks[-1].page_start is None


def test_docx_inherited_heading_nested_lists_hyperlinks_and_merged_tables(settings):
    document = Document()
    custom = document.styles.add_style("Custom section", WD_STYLE_TYPE.PARAGRAPH)
    custom.base_style = document.styles["Heading 2"]
    document.add_heading("Root", 1)
    paragraph = document.add_paragraph("Inherited heading", style=custom)
    paragraph = document.add_paragraph("Linked ")
    hyperlink = OxmlElement("w:hyperlink")
    run, text = OxmlElement("w:r"), OxmlElement("w:t")
    text.text = "reference"
    run.append(text)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)
    paragraph = document.add_paragraph("Nested item", style="List Bullet")
    numbering = paragraph._p.get_or_add_pPr().get_or_add_numPr()
    numbering.get_or_add_numId().val = 1
    numbering.get_or_add_ilvl().val = 1
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "Shared"
    table.cell(1, 0).text = "Multi\nline"
    nested = table.cell(1, 1).add_table(rows=1, cols=1)
    nested.cell(0, 0).text = "Nested cell"
    output = BytesIO()
    document.save(output)
    parsed = parse_binary(output.getvalue(), "docx", settings)
    assert parsed.blocks[1].heading_path == ["Root", "Inherited heading"]
    assert parsed.blocks[2].text == "Linked reference"
    assert parsed.blocks[3].level == 2 and "Nested item" in parsed.blocks[3].text
    assert parsed.blocks[-1].text.count("Shared") == 1
    assert "Multi / line" in parsed.blocks[-1].text
    assert "Nested cell" in parsed.blocks[-1].text


@pytest.mark.parametrize("data", [b"%PDF-1.7\nSECRET corrupt", binary_fixture("pdf")[:500]])
def test_malformed_pdf(settings, data):
    assert_error(data, "pdf", settings, "invalid_pdf")


def test_scanned_empty_near_empty_and_encrypted_pdf(settings):
    assert_error(
        (FIXTURES / "image-only.pdf").read_bytes(), "pdf", settings, "pdf_text_unavailable"
    )
    writer = PdfWriter()
    writer.add_blank_page(100, 100)
    output = BytesIO()
    writer.write(output)
    assert_error(output.getvalue(), "pdf", settings, "pdf_text_unavailable")
    # Real digital text can also be insufficient under a configured threshold.
    assert_error(
        binary_fixture("pdf"),
        "pdf",
        settings.model_copy(update={"min_pdf_text_chars": 1000}),
        "pdf_text_unavailable",
    )
    writer = PdfWriter(clone_from=BytesIO(binary_fixture("pdf")))
    writer.encrypt("password")
    output = BytesIO()
    writer.write(output)
    assert_error(output.getvalue(), "pdf", settings, "pdf_encrypted")


@pytest.mark.parametrize(
    "updates", [{"max_pdf_pages": 2}, {"max_document_objects": 3}, {"max_decompressed_bytes": 10}]
)
def test_pdf_resource_bounds(settings, updates):
    assert_error(
        binary_fixture("pdf"), "pdf", settings.model_copy(update=updates), "pdf_resource_limit"
    )


def test_pdf_compressed_stream_bound(settings):
    reader = PdfReader(BytesIO(binary_fixture("pdf")))
    writer = PdfWriter()
    page = writer.add_page(reader.pages[0])
    content = DecodedStreamObject()
    content.set_data(b" " * 100_000 + page.get_contents().get_data())
    page[NameObject("/Contents")] = writer._add_object(content.flate_encode())
    output = BytesIO()
    writer.write(output)
    assert len(output.getvalue()) < 10_000
    assert_error(
        output.getvalue(),
        "pdf",
        settings.model_copy(update={"max_pdf_stream_bytes": 1024}),
        "pdf_resource_limit",
    )


@pytest.mark.parametrize("data", [b"PK\x03\x04SECRET corrupt", binary_fixture("docx")[:500]])
def test_malformed_docx(settings, data):
    assert_error(data, "docx", settings, "invalid_docx")


@pytest.mark.parametrize(
    "updates",
    [
        {"max_archive_members": 2},
        {"max_decompressed_bytes": 100},
        {"max_compression_ratio": 1},
        {"max_document_objects": 10},
        {"max_parse_depth": 2},
    ],
)
def test_docx_resource_bounds(settings, updates):
    assert_error(
        binary_fixture("docx"), "docx", settings.model_copy(update=updates), "docx_resource_limit"
    )


def test_docx_zip_bomb_and_xml_entities(settings):
    bomb = repack_docx(additions={"unused.txt": b"x" * 2_000_000})
    assert len(bomb) < 50_000
    assert_error(bomb, "docx", settings, "docx_resource_limit")
    for encoding in ("utf-8", "utf-16"):
        xml = (
            '<?xml version="1.0" encoding="' + encoding + '"?>'
            '<!DOCTYPE foo [<!ENTITY secret SYSTEM "file:///etc/passwd">]><foo>&secret;</foo>'
        )
        hostile = repack_docx(additions={"unused.xml": xml.encode(encoding)})
        assert_error(hostile, "docx", settings, "invalid_docx")


@pytest.mark.parametrize("name", ["../escape.xml", "/absolute.xml", "word/vbaProject.bin"])
def test_docx_unsafe_members(settings, name):
    assert_error(repack_docx(additions={name: b"data"}), "docx", settings, "invalid_docx")


def test_docx_duplicate_entries_invalid_xml_grid_abuse_and_empty(settings):
    with pytest.warns(UserWarning):
        duplicate = repack_docx(additions={"word/document.xml": b"SECRET"})
    assert_error(duplicate, "docx", settings, "invalid_docx")
    invalid = repack_docx(changes={"word/document.xml": lambda _: b"<broken>SECRET"})
    assert_error(invalid, "docx", settings, "invalid_docx")
    grid = repack_docx(
        changes={
            "word/document.xml": lambda value: value.replace(
                b"<w:tcPr>", b'<w:tcPr><w:gridSpan w:val="999999999"/>', 1
            )
        }
    )
    assert_error(grid, "docx", settings, "docx_resource_limit")
    empty = BytesIO()
    Document().save(empty)
    assert_error(empty.getvalue(), "docx", settings, "docx_text_unavailable")


@pytest.mark.parametrize("fmt", ["pdf", "docx"])
def test_binary_input_output_and_block_limits(settings, fmt):
    data = binary_fixture(fmt)
    assert_error(data, fmt, settings.model_copy(update={"max_file_bytes": 10}), "input_too_large")
    assert_error(
        data, fmt, settings.model_copy(update={"max_normalized_bytes": 100}), "normalized_too_large"
    )
    # Separate from preflight object/node budgets.
    assert_error(
        data, fmt, settings.model_copy(update={"max_parse_nodes": 1}), "structure_too_large"
    )


@pytest.mark.parametrize(
    "data,name,mime,override,expected",
    [
        (binary_fixture("pdf"), "../unknown.dat", "application/octet-stream", "auto", "pdf"),
        (binary_fixture("docx"), "WORD.DOCX", None, "auto", "docx"),
        (b"# plain", "fake.json", "application/json", "text", "text"),
    ],
)
def test_signature_selection(data, name, mime, override, expected):
    assert select_upload_format(data, name, mime, override) == expected


@pytest.mark.parametrize(
    "data,name,mime,override",
    [
        (b"plain text", "fake.pdf", None, "auto"),
        (binary_fixture("pdf"), "fake.docx", None, "auto"),
        (binary_fixture("docx"), "fake.dat", "application/pdf", "auto"),
        (binary_fixture("docx"), "fake.dat", "Application/PDF; charset=utf-8", "auto"),
        (binary_fixture("pdf"), "actual.pdf", None, "text"),
        (b"plain text", "fake.txt", None, "docx"),
    ],
)
def test_conflicting_binary_hints(data, name, mime, override):
    with pytest.raises(ParseError) as caught:
        select_upload_format(data, name, mime, override)
    assert caught.value.code == "file_type_mismatch"


def test_subprocess_timeout_cleanup_and_no_secrets(settings, monkeypatch):
    from knowledge_bootstrap import binary

    original_run = subprocess.run
    paths = []

    def slow_child(command, **kwargs):
        limits = json.loads(command[4])
        assert "database_url" not in limits and "auth_tokens" not in limits
        assert not any(key.startswith("KNOWLEDGE_") for key in kwargs["env"])
        paths.append(Path(command[-1]).parent)
        return original_run([sys.executable, "-c", "import time; time.sleep(10)"], **kwargs)

    monkeypatch.setattr(binary.subprocess, "run", slow_child)
    assert_error(
        binary_fixture("pdf"),
        "pdf",
        settings.model_copy(update={"parser_timeout_seconds": 0.1}),
        "parser_timeout",
    )
    assert paths and all(not path.exists() for path in paths)


def test_subprocess_crash_is_stable_and_cleans_up(settings, monkeypatch):
    from knowledge_bootstrap import binary

    paths = []

    def killed_child(command, **kwargs):
        paths.append(Path(command[-1]).parent)
        return subprocess.CompletedProcess(command, -9)

    monkeypatch.setattr(binary.subprocess, "run", killed_child)
    assert_error(binary_fixture("pdf"), "pdf", settings, "parser_resource_limit")
    assert all(not path.exists() for path in paths)


def test_docx_cyclic_styles(settings):
    document = Document()
    a = document.styles.add_style("A", WD_STYLE_TYPE.PARAGRAPH)
    b = document.styles.add_style("B", WD_STYLE_TYPE.PARAGRAPH)
    a.base_style, b.base_style = b, a
    document.add_paragraph("SECRET", style=a)
    output = BytesIO()
    document.save(output)
    assert_error(output.getvalue(), "docx", settings, "invalid_docx")
