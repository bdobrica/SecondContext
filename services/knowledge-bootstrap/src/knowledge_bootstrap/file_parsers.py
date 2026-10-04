"""Digital PDF and DOCX extraction. Called only by the bounded subprocess."""

import re
from io import BytesIO
from pathlib import PurePosixPath
from zipfile import ZIP_DEFLATED, ZIP_STORED, BadZipFile, ZipFile

from defusedxml.common import DefusedXmlException
from defusedxml.ElementTree import iterparse
from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph
from pypdf import PdfReader, apply_configuration
from pypdf.errors import LimitReachedError

from knowledge_bootstrap.config import Settings
from knowledge_bootstrap.parsers import Block, ParsedDocument, ParseError, validate_document


def _clean_text(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    if any((ord(c) < 32 and c not in "\n\t") or 127 <= ord(c) < 160 for c in text):
        raise ParseError("invalid_extracted_text", "Extracted text contains unsupported controls")
    return "\n".join(line.rstrip() for line in text.split("\n")).strip()


def _pdf(data: bytes, settings: Settings) -> ParsedDocument:
    stream_limit = settings.max_pdf_stream_bytes
    with apply_configuration(
        maximum_declared_stream_length=stream_limit,
        array_based_stream_maximum_output_length=stream_limit,
        zlib_maximum_output_length=stream_limit,
        lzw_maximum_output_length=stream_limit,
        run_length_maximum_output_length=stream_limit,
        flate_maximum_row_length=stream_limit,
        image_maximum_buffer_size=stream_limit,
        page_tree_maximum_entries=settings.max_document_objects,
        page_tree_maximum_depth=settings.max_parse_depth,
        xform_maximum_invocations_per_extraction=settings.max_document_objects,
        jbig2dec_binary=None,
    ):
        reader = PdfReader(BytesIO(data), strict=True)
        if reader.is_encrypted:
            raise ParseError(
                "pdf_encrypted", "Encrypted PDFs are unsupported; upload a decrypted copy"
            )
        objects = sum(len(entries) for entries in reader.xref.values()) + len(reader.xref_objStm)
        if objects > settings.max_document_objects:
            raise ParseError("pdf_resource_limit", "PDF exceeds the object limit")
        count = len(reader.pages)
        if count > settings.max_pdf_pages:
            raise ParseError("pdf_resource_limit", "PDF exceeds the page limit")
        blocks, pages, low_text = [], [], []
        decompressed = 0
        text_bytes = 0
        for number, page in enumerate(reader.pages, start=1):
            contents = page.get_contents()
            if contents is not None:
                size = len(contents.get_data())
                decompressed += size
                if size > stream_limit or decompressed > settings.max_decompressed_bytes:
                    raise ParseError("pdf_resource_limit", "PDF exceeds the decoded content limit")
            text = _clean_text(
                page.extract_text(extraction_mode="layout", layout_mode_strip_rotated=False)
                if contents is not None
                else ""
            )
            text_bytes += len(text.encode())
            if text_bytes > settings.max_normalized_bytes:
                raise ParseError(
                    "normalized_too_large", "Extracted PDF exceeds the output byte limit"
                )
            if sum(c.isalnum() for c in text) < settings.min_pdf_text_chars:
                low_text.append(number)
            # Layout extraction retains vertical gaps; a gap is a paragraph boundary.
            for paragraph in re.split(r"\n[ \t]*\n+", text):
                if paragraph.strip():
                    blocks.append(
                        Block("paragraph", paragraph.strip(), page_start=number, page_end=number)
                    )
            if len(blocks) > settings.max_parse_nodes:
                raise ParseError("structure_too_large", "PDF exceeds the block limit")
            pages.append(text)
        text = "\n\n".join(page for page in pages if page)
        if sum(c.isalnum() for c in text) < settings.min_pdf_text_chars or len(low_text) == count:
            raise ParseError(
                "pdf_text_unavailable",
                "PDF has insufficient digital text; scanned/image-only or empty PDFs need OCR, "
                "which is unsupported",
            )
        return ParsedDocument(
            "pdf",
            text,
            blocks,
            extra_metadata={"page_count": count, "pages_with_low_text": low_text},
        )


def _preflight_docx(data: bytes, settings: Settings) -> None:
    """Check declared and actual expansion, CRCs and XML before python-docx loads it.

    No member is extracted to disk. Defused XML forbids DTD/entities/external references,
    even in UTF-16 XML. All XML/relationship parts are checked, including unused parts.
    """
    with ZipFile(BytesIO(data)) as archive:
        members = archive.infolist()
        if len(members) > settings.max_archive_members:
            raise ParseError("docx_resource_limit", "DOCX exceeds the ZIP member limit")
        names = [member.filename for member in members]
        if len(names) != len(set(names)) or not {"[Content_Types].xml", "word/document.xml"} <= set(
            names
        ):
            raise ParseError("invalid_docx", "ZIP is not an unambiguous DOCX package")
        declared = 0
        for member in members:
            path = PurePosixPath(member.filename)
            if path.is_absolute() or ".." in path.parts or "\\" in member.filename:
                raise ParseError("invalid_docx", "DOCX contains an invalid member path")
            if member.flag_bits & 1 or member.compress_type not in {ZIP_STORED, ZIP_DEFLATED}:
                raise ParseError("invalid_docx", "Unsupported ZIP encryption or compression")
            if member.filename.lower().endswith("vbaproject.bin"):
                raise ParseError("invalid_docx", "Macro-enabled documents are unsupported")
            declared += member.file_size
            if (
                declared > settings.max_decompressed_bytes
                or member.file_size > max(1, member.compress_size) * settings.max_compression_ratio
            ):
                raise ParseError("docx_resource_limit", "DOCX exceeds ZIP expansion limits")
        expanded, nodes = 0, 0
        for member in members:
            with archive.open(member) as stream:
                content = stream.read(settings.max_decompressed_bytes - expanded + 1)
            expanded += len(content)
            if expanded > settings.max_decompressed_bytes:
                raise ParseError("docx_resource_limit", "DOCX exceeds ZIP expansion limits")
            if not member.filename.lower().endswith((".xml", ".rels")):
                continue
            depth = 0
            for event, element in iterparse(
                BytesIO(content), events=("start", "end"), forbid_dtd=True
            ):
                if event == "start":
                    depth += 1
                    nodes += 1
                    if depth > settings.max_parse_depth or nodes > settings.max_document_objects:
                        raise ParseError("docx_resource_limit", "DOCX exceeds XML structure limits")
                    # Word grid spans can expand a few XML nodes into huge Python cell lists.
                    if element.tag in {qn("w:gridSpan"), qn("w:gridBefore"), qn("w:gridAfter")}:
                        value = element.get(qn("w:val"), "1")
                        if (
                            len(value) > 6
                            or not value.isdigit()
                            or int(value) > settings.max_parse_nodes
                        ):
                            raise ParseError(
                                "docx_resource_limit", "DOCX exceeds table grid limits"
                            )
                else:
                    depth -= 1
                    element.clear()


def _properties(paragraph: Paragraph, settings: Settings):
    if paragraph._p.pPr is not None:
        yield paragraph._p.pPr
    style, seen = paragraph.style, set()
    while style is not None:
        if style.style_id in seen or len(seen) >= settings.max_parse_depth:
            raise ParseError("invalid_docx", "DOCX contains cyclic or excessive style inheritance")
        seen.add(style.style_id)
        if style.element.pPr is not None:
            yield style.element.pPr
        style = style.base_style


def _paragraph_structure(paragraph: Paragraph, settings: Settings):
    outline, numbering = None, None
    for properties in _properties(paragraph, settings):
        if outline is None and (element := properties.find(qn("w:outlineLvl"))) is not None:
            outline = int(element.get(qn("w:val")))
        if numbering is None and (element := properties.find(qn("w:numPr"))) is not None:
            num = element.find(qn("w:numId"))
            level = element.find(qn("w:ilvl"))
            if num is not None:
                numbering = (
                    num.get(qn("w:val")),
                    int(level.get(qn("w:val"))) if level is not None else 0,
                )
    if outline is not None:
        heading = outline + 1 if 0 <= outline <= 8 else None
    else:
        match = re.fullmatch(r"Heading ([1-9])", paragraph.style.name or "", re.IGNORECASE)
        heading = int(match[1]) if match else None
    if numbering and (numbering[0] == "0" or not 0 <= numbering[1] <= 8):
        numbering = None
    return heading, numbering


def _list_kind(document, num_id: str, level: int) -> str:
    numbering = document.part.numbering_part.element
    num = next((n for n in numbering.findall(qn("w:num")) if n.get(qn("w:numId")) == num_id), None)
    if num is not None:
        abstract_id = num.find(qn("w:abstractNumId")).get(qn("w:val"))
        for abstract in numbering.findall(qn("w:abstractNum")):
            if abstract.get(qn("w:abstractNumId")) == abstract_id:
                for definition in abstract.findall(qn("w:lvl")):
                    if definition.get(qn("w:ilvl")) == str(level):
                        fmt = definition.find(qn("w:numFmt"))
                        return (
                            "bullet"
                            if fmt is not None and fmt.get(qn("w:val")) == "bullet"
                            else "number"
                        )
    return "number"


def _table_text(table: Table, settings: Settings, depth: int = 1) -> str:
    if depth > settings.max_parse_depth:
        raise ParseError("docx_resource_limit", "DOCX exceeds table nesting limits")
    rows = []
    seen = set()
    cells = 0
    for row in table.rows:
        values = []
        for cell in row.cells:
            cells += 1
            if cells > settings.max_parse_nodes:
                raise ParseError("docx_resource_limit", "DOCX exceeds the table cell limit")
            if cell._tc in seen:
                values.append("")  # Keep grid position without repeating merged-cell text.
                continue
            seen.add(cell._tc)
            parts = [
                item.text if isinstance(item, Paragraph) else _table_text(item, settings, depth + 1)
                for item in cell.iter_inner_content()
            ]
            value = _clean_text("\n".join(parts)).replace("|", "\\|").replace("\n", " / ")
            values.append(value)
        rows.append(" | ".join(values))
    return "\n".join(rows)


def _docx(data: bytes, settings: Settings) -> ParsedDocument:
    _preflight_docx(data, settings)
    document = Document(BytesIO(data))
    blocks, headings, counters = [], [], {}
    title = None
    total = 0
    for item in document.iter_inner_content():
        level, numbering = (
            _paragraph_structure(item, settings) if isinstance(item, Paragraph) else (None, None)
        )
        text = _clean_text(
            item.text if isinstance(item, Paragraph) else _table_text(item, settings)
        )
        if not text:
            continue
        kind = "table" if isinstance(item, Table) else "paragraph"
        if level is not None:
            while headings and headings[-1][0] >= level:
                headings.pop()
            headings.append((level, text))
            title = title or text
            kind = "heading"
        elif numbering:
            num_id, list_level = numbering
            key = (num_id, list_level)
            counters[key] = counters.get(key, 0) + 1
            for deeper in list(counters):
                if deeper[0] == num_id and deeper[1] > list_level:
                    del counters[deeper]
            marker = (
                "-" if _list_kind(document, num_id, list_level) == "bullet" else f"{counters[key]}."
            )
            text = "  " * list_level + marker + " " + text
            kind, level = "list", list_level + 1
        elif isinstance(item, Paragraph) and item.style.name == "Title":
            title = title or text
        total += len(text.encode())
        if total > settings.max_normalized_bytes:
            raise ParseError("normalized_too_large", "Extracted DOCX exceeds the output byte limit")
        blocks.append(Block(kind, text, [label for _, label in headings], level=level))
        if len(blocks) > settings.max_parse_nodes:
            raise ParseError("structure_too_large", "DOCX exceeds the block limit")
    if not blocks:
        raise ParseError("docx_text_unavailable", "DOCX has no supported body text")
    return ParsedDocument("docx", "\n\n".join(block.text for block in blocks), blocks, title)


def parse_file(data: bytes, fmt: str, settings: Settings) -> ParsedDocument:
    # Quiet third-party diagnostics: malformed content must not leak into logs/results.
    try:
        document = _pdf(data, settings) if fmt == "pdf" else _docx(data, settings)
        return validate_document(document, settings)
    except ParseError:
        raise
    except LimitReachedError:
        raise ParseError("pdf_resource_limit", "PDF exceeds parser resource limits") from None
    except (BadZipFile, DefusedXmlException):
        raise ParseError("invalid_docx", "Malformed or unsafe DOCX package") from None
    except MemoryError:
        raise
    except Exception:
        raise ParseError("invalid_" + fmt, "Malformed or unsupported file") from None
